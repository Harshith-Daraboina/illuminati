# Required libraries: pandas, openpyxl
# Install them if you don't have them: pip install pandas openpyxl

import os
import pandas as pd
from datetime import datetime, time, timedelta
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill
from openpyxl.utils import get_column_letter
import logging
import random
import math
from collections import defaultdict
import copy # Needed for deep copying schedule states

# Set up logging
logging.basicConfig(level=logging.INFO, format='%(asctime)s - %(levelname)s - %(message)s')

class TimeTableConfig:
    """Configuration settings for timetable generation"""
    DAYS = ['Monday', 'Tuesday', 'Wednesday', 'Thursday', 'Friday']
    START_TIME = time(9, 0)
    END_TIME = time(18, 30)
    
    # Break times
    LUNCH_BREAK_START = time(13, 0)
    LUNCH_BREAK_END = time(14, 30)
    MORNING_BREAK_START = time(10, 30)
    MORNING_BREAK_END = time(11, 0)
    
    # Slot durations
    SLOT_DURATION_MINUTES = 30
    LECTURE_SLOTS = 3
    LAB_SLOTS = 4
    TUTORIAL_SLOTS = 2
    
    # Scheduling attempts
    MAX_SCHEDULING_ATTEMPTS_BASKET = 300
    MAX_SCHEDULING_ATTEMPTS_SESSION = 500
    
    # File paths
    CLASSROOMS_FILE = 'Classrooms.csv'
    COURSES_FILE = 'Semester_courses.csv'
    OUTPUT_FILE = 'timetable_output.xlsx'
    
    # Lab types by department
    LAB_TYPES = {
        'CSE': 'COMPUTER_LAB',
        'DSAI': 'COMPUTER_LAB',
        'ECE': 'HARDWARE_LAB'
    }

# --- Time Slot Generation ---
def generate_time_slots():
    """Generate all possible 30-min time slots in a day, identifying breaks."""
    slots = [] # List of all slot details including breaks
    schedulable_slot_indices = [] # List of indices of slots that are NOT breaks
    current_time = datetime.combine(datetime.today(), TimeTableConfig.START_TIME)
    # Calculate end_datetime to include the slot starting exactly at END_TIME if SLOT_DURATION_MINUTES allows
    end_limit_datetime = datetime.combine(datetime.today(), TimeTableConfig.END_TIME) + timedelta(minutes=TimeTableConfig.SLOT_DURATION_MINUTES)

    slot_index = 0
    while current_time < end_limit_datetime:
        start = current_time.time()
        end = (current_time + timedelta(minutes=TimeTableConfig.SLOT_DURATION_MINUTES)).time()
        # Cap the end time at END_TIME if it goes beyond
        if end > TimeTableConfig.END_TIME:
            end = TimeTableConfig.END_TIME

        is_break = False
        # A slot is a break slot if its START time is within a defined break period
        if (TimeTableConfig.MORNING_BREAK_START <= start < TimeTableConfig.MORNING_BREAK_END) or \
           (TimeTableConfig.LUNCH_BREAK_START <= start < TimeTableConfig.LUNCH_BREAK_END):
            is_break = True

        slot_label = f"{start.strftime('%H:%M')}-{end.strftime('%H:%M')}"
        slot_info = {
           "index": slot_index,
           "label": slot_label,
           "start_time": start,
           "end_time": end,
           "is_break": is_break
        }
        slots.append(slot_info)

        if not is_break:
            schedulable_slot_indices.append(slot_index) # Store indices of non-break slots

        current_time += timedelta(minutes=TimeTableConfig.SLOT_DURATION_MINUTES)
        slot_index += 1

    num_slots_per_day = len(slots)
    logging.info(f"Generated {num_slots_per_day} total slots per day.")
    logging.info(f"Identified {len(schedulable_slot_indices)} schedulable slot indices.")

    return slots, schedulable_slot_indices, num_slots_per_day

ALL_SLOTS, SCHEDULABLE_SLOT_INDICES, NUM_SLOTS_PER_DAY = generate_time_slots()


# --- Data Loading ---
def load_classroom_data(file_path=TimeTableConfig.CLASSROOMS_FILE):
    """Load classroom data from CSV file, add priority."""
    if not os.path.exists(file_path):
        logging.error(f"Classrooms data not found at {file_path}")
        raise FileNotFoundError(f"Classrooms data not found at {file_path}")

    try:
        df = pd.read_csv(file_path)
        df = df.dropna(subset=['Classroom']) # Remove rows where Classroom is NaN
        df['Capacity'] = pd.to_numeric(df['Capacity'], errors='coerce').fillna(0).astype(int)
        df['Type'] = df['Type'].astype(str).str.upper() # Standardize type

        # Prioritize larger rooms and specific lecture halls
        df['Priority'] = df.apply(
             lambda row: 1 if row['Capacity'] >= 100 and 'LECTURE' in row['Type']
             else 2 if 'LECTURE' in row['Type']
             else 3 if ('COMPUTER_LAB' in row['Type'] or 'HARDWARE_LAB' in row['Type']) and row['Capacity'] >= 40 # Prioritize larger labs slightly
             else 4 if ('COMPUTER_LAB' in row['Type'] or 'HARDWARE_LAB' in row['Type'])
             else 5, axis=1 # Other types lower priority
        )
        # Sort by priority (lower is better), then capacity (higher is better)
        return df.sort_values(['Priority', 'Capacity'], ascending=[True, False]).to_dict('records')
    except Exception as e:
        logging.error(f"Error loading or processing classroom data: {e}")
        raise

def load_course_data(file_path=TimeTableConfig.COURSES_FILE):
    """Load course data from CSV file, clean and calculate needed slots."""
    if not os.path.exists(file_path):
        logging.error(f"Course data not found at {file_path}")
        raise FileNotFoundError(f"Course data not found at {file_path}")

    try:
        df = pd.read_csv(file_path)
        # Drop rows where essential info is missing
        df = df.dropna(subset=['Department', 'Semester', 'Course Code', 'Course Name', 'Faculty'])
        df['Semester'] = df['Semester'].astype(str).str.strip()
        df['Department'] = df['Department'].astype(str).str.strip()
        df['Course Code'] = df['Course Code'].astype(str).str.strip()
        df['Course Name'] = df['Course Name'].astype(str).str.strip()
        df['Faculty'] = df['Faculty'].astype(str).str.strip()

        df['Enrolled_students'] = pd.to_numeric(df['Enrolled_students'], errors='coerce').fillna(25).astype(int) # Default enrollment if missing
        df['L'] = pd.to_numeric(df['L'], errors='coerce').fillna(0)
        df['T'] = pd.to_numeric(df['T'], errors='coerce').fillna(0)
        df['P'] = pd.to_numeric(df['P'], errors='coerce').fillna(0)
        df['class_connector'] = df['class_connector'].astype(str).replace('nan', '').str.strip()


        # Identify electives (Basket courses) - check Course Code for B1/B2 etc.
        df['is_elective'] = df['Course Code'].str.contains(r'^B[1-4]', regex=True) | df['Course Name'].str.contains(r'^B[1-4]', regex=True) # Check code or name
        df['basket_code'] = df['Course Code'].str.extract(r'^(B[1-4])', expand=False).fillna('')
        # If basket code not in Course Code, try Course Name (e.g., "B1(ASD151/...)")
        mask = (df['basket_code'] == '') & df['is_elective']
        df.loc[mask, 'basket_code'] = df.loc[mask, 'Course Name'].str.extract(r'^(B[1-4])', expand=False).fillna('')

        # Ensure all electives have a basket code, default if necessary
        df.loc[df['is_elective'] & (df['basket_code'] == ''), 'basket_code'] = 'B_Unknown'


        # Calculate required weekly sessions (each session is a contiguous block of slots)
        df['L_sessions_needed'] = df['L'].apply(lambda x: math.ceil(x / (TimeTableConfig.LECTURE_SLOTS * TimeTableConfig.SLOT_DURATION_MINUTES / 60)) if TimeTableConfig.LECTURE_SLOTS > 0 else 0)
        df['T_sessions_needed'] = df['T'].apply(lambda x: math.ceil(x / (TimeTableConfig.TUTORIAL_SLOTS * TimeTableConfig.SLOT_DURATION_MINUTES / 60)) if TimeTableConfig.TUTORIAL_SLOTS > 0 else 0)
        # P sessions need 2 rooms per session
        df['P_sessions_needed'] = df['P'].apply(lambda x: math.ceil(x / (TimeTableConfig.LAB_SLOTS * TimeTableConfig.SLOT_DURATION_MINUTES / 60)) if TimeTableConfig.LAB_SLOTS > 0 else 0)

        # Add columns to track scheduled sessions per course component (for DataFrame status)
        df['L_scheduled'] = 0
        df['T_scheduled'] = 0
        df['P_scheduled'] = 0
        df['failed'] = [[] for _ in range(len(df))] # To store failed session types (list per row)


        # Unique identifier for each course instance
        df['unique_id'] = df.apply(lambda row: f"{row['Department']}_{row['Semester']}_{row['Course Code']}_{row['Faculty']}", axis=1)
        # Use Department, Semester, Course Code, and Faculty for uniqueness


        return df
    except Exception as e:
        logging.error(f"Error loading or processing course data: {e}")
        raise


# --- Timetable Initialization ---

def initialize_timetable(departments, semesters_by_dept):
    """Create empty timetable structures for all department-semester combinations."""
    timetables = {}
    for dept in departments:
        # Ensure semesters_by_dept[dept] is iterable
        dept_semesters = semesters_by_dept.get(dept, [])
        if not hasattr(dept_semesters, '__iter__'):
            logging.warning(f"Semesters for department {dept} not found or not iterable. Skipping.")
            continue

        for sem in dept_semesters:
            key = f"{dept}_{sem}"
            # Create a 2D list: [day][slot]
            # Initialize with default empty slot structure, including a list for rooms
            timetables[key] = [
                [
                    {'type': None, 'course': None, 'code': None, 'rooms': [], 'prof': None, 'span': 0, 'connector': None, 'basket': None}
                    for _ in range(NUM_SLOTS_PER_DAY) # Use the calculated total number of slots
                ]
                for _ in range(len(TimeTableConfig.DAYS)) # For each day of the week
            ]
            logging.debug(f"Initialized timetable structure for {key}")
    return timetables

# --- Scheduling Helper Functions ---

def get_lab_type(department):
    """Return lab type based on department."""
    return TimeTableConfig.LAB_TYPES.get(department.upper(), 'COMPUTER_LAB')

def is_slot_range_valid(start_slot_index, num_slots):
    """Check if a range of slots is valid (within day, no breaks)."""
    if start_slot_index < 0 or start_slot_index + num_slots > NUM_SLOTS_PER_DAY:
        return False
    for i in range(num_slots):
        if ALL_SLOTS[start_slot_index + i]['is_break']:
            return False
    return True

def check_availability(schedule, key, day_index, start_slot_index, num_slots):
    """Check if a resource (prof, room, group) is free."""
    # Check if the primary key exists in the schedule
    if key not in schedule:
        return True # Resource not yet scheduled, so available

    # Check if the day_index exists for this resource
    if day_index not in schedule[key]:
        return True # Resource not scheduled on this day, so available

    day_schedule_slots = schedule[key][day_index] # This should be a set of busy slot indices

    for i in range(num_slots):
        if (start_slot_index + i) in day_schedule_slots:
            return False # Slot is busy
    return True # All slots in range are free

def book_slot(schedule, key, day_index, start_slot_index, num_slots):
    """Mark slots as busy for a resource."""
    if key not in schedule:
        schedule[key] = {}
    if day_index not in schedule[key]:
        schedule[key][day_index] = set() # Use a set for efficient lookup

    for i in range(num_slots):
        schedule[key][day_index].add(start_slot_index + i)


def find_available_rooms(classrooms, required_type, required_capacity, day_index, start_slot_index, num_slots, classroom_schedule, num_rooms=1, excluded_rooms=None):
    """Find suitable and available room(s), excluding specified rooms."""
    if excluded_rooms is None: excluded_rooms = set()
    found_rooms = []
    current_excluded = set(excluded_rooms) # Copy to add rooms found within this call

    potential_rooms = [
        r for r in classrooms
        if required_type in r['Type'].upper()
        and r['Capacity'] >= required_capacity
    ]

    # Sort potential rooms by priority and capacity to prefer better rooms first
    potential_rooms.sort(key=lambda r: (r['Priority'], -r['Capacity']))

    for room in potential_rooms:
        room_id = room['Classroom']
        if room_id not in current_excluded:
            # Check availability for this room in the classroom_schedule state
            if check_availability(classroom_schedule, room_id, day_index, start_slot_index, num_slots):
                found_rooms.append(room_id)
                current_excluded.add(room_id) # Exclude this room from subsequent searches in this call
                if len(found_rooms) == num_rooms:
                    return found_rooms # Found all required rooms

    return [] # Did not find enough rooms


# --- Connector and Basket Grouping ---

def build_connector_groups(courses_df):
    """Group courses by class_connector, calculate combined needs."""
    connector_groups = defaultdict(lambda: {
        'connector_id': None, # Explicitly store the connector ID
        'courses': [], # List of original course data dicts
        'total_enrollment': 0,
        'professor': None, # Assumes one professor for the group
        'participating_groups': set(), # Store dept_sem keys
        'P_sessions_scheduled': 0,
        'failed_sessions': [], # e.g., [{'type': 'L', 'count': 1}]
        'is_elective_connector': False, # Track if it's for electives
        'basket_code': None # Store basket code if applicable
    })

    for _, course in courses_df.iterrows():
        connector_id = course['class_connector']
        if connector_id:
            group = connector_groups[connector_id]
            group['connector_id'] = connector_id # Set the ID

            course_info = course.to_dict()
            group['courses'].append(course_info)
            group['total_enrollment'] = max(group['total_enrollment'], course['Enrolled_students']) # Use max enrollment

            # For simplicity, use the needs from the *first* course added to the group as the group's requirement.
            # This might need refinement if courses within a connector have different L/T/P total hour requirements.
            if len(group['courses']) == 1: # First course being added
                group['L_sessions_needed'] = course['L_sessions_needed']
                group['T_sessions_needed'] = course['T_sessions_needed']
                group['P_sessions_needed'] = course['P_sessions_needed']

            group['participating_groups'].add(f"{course['Department']}_{course['Semester']}")

            if course['is_elective']:
                group['is_elective_connector'] = True # Mark if any course in group is elective
                group['basket_code'] = course['basket_code'] # Store basket code

    # Convert defaultdict to dict for easier processing later
    return dict(connector_groups)


def group_courses_by_basket(courses_df, connector_groups):
    """Group elective courses by basket, identifying connector groups and standalone courses within."""
    baskets = defaultdict(lambda: {
        'basket_code': None,
        'semester_logic_group': None, # e.g., "Sem6" or "Sem2"
        'participating_groups': set(), # Dept_Sem keys of groups who can take this basket
        'offerings': [], # List of distinct offerings (standalone course dict or connector group key)
        'needed_duration': LECTURE_SLOTS, # Assume the common basket time is for the Lecture component
        'scheduled': False,
        'scheduled_day': None,
        'scheduled_start_slot': None,
        'scheduled_assignments': {} # offering_key -> {'prof': prof, 'room': room_id}
    })

    # Determine logical semester groups for baskets (e.g., 2A/2B -> Sem2)
    def get_semester_group(sem):
        if sem.startswith('2'): return 'Sem2'
        if sem.startswith('4'): return 'Sem4'
        if sem.startswith('6'): return 'Sem6'
        if sem.startswith('8'): return 'Sem8'
        return sem # Fallback

    for _, course in courses_df[courses_df['is_elective']].iterrows():
        basket_code = course['basket_code']
        if not basket_code or basket_code == 'B_Unknown': continue # Skip if not properly identified as basket

        semester_group = get_semester_group(course['Semester'])
        basket_key = f"{semester_group}_{basket_code}" # e.g., Sem6_B1

        group = baskets[basket_key]
        group['basket_code'] = basket_code
        group['semester_logic_group'] = semester_group
        group['participating_groups'].add(f"{course['Department']}_{course['Semester']}")

        connector_id = course['class_connector']
        # Offering key is the connector ID if it's an elective connector, otherwise the unique course ID (for standalone electives)
        offering_key = connector_id if connector_id and connector_groups.get(connector_id, {}).get('is_elective_connector', False) else course['unique_id']

        # Add offering only once per basket group using the offering_key
        # Keep track of offering keys already added to avoid duplicates
        existing_offering_keys = [o if isinstance(o, str) else o['unique_id'] for o in group['offerings']]

        if offering_key not in existing_offering_keys:
            if offering_key == connector_id: # It's an elective connector group
                group['offerings'].append(connector_id)
            else: # It's a standalone elective course
                group['offerings'].append(course.to_dict())

            # The basket block assumes LECTURE_SLOTS duration for the common time.
            # Individual T/P sessions for electives need to be scheduled separately.

    return dict(baskets)


# --- Scheduling Functions ---

def schedule_basket_group(basket_key, basket_info, classrooms, connector_groups,
                         professor_schedule, classroom_schedule, group_schedule, timetables, course_schedule_status_df, scheduled_sessions_list):
    """Attempt to schedule all offerings within a basket simultaneously in distinct rooms."""
    if basket_info['scheduled']: return True # Already done

    logging.info(f"Attempting to schedule Basket Group: {basket_key}")

    duration = basket_info['needed_duration'] # Assumes LECTURE_SLOTS for the common block
    offerings = basket_info['offerings']
    participating_dept_sem_groups = basket_info['participating_groups']

    attempts = 0
    scheduled = False
    while not scheduled and attempts < MAX_SCHEDULING_ATTEMPTS_BASKET:
        attempts += 1
        day_index = random.randrange(len(DAYS))
        # Pick a random valid starting slot index from schedulable slots
        if not SCHEDULABLE_SLOT_INDICES:
            logging.error("No schedulable slots available!")
            return False
        start_slot_index = random.choice(SCHEDULABLE_SLOT_INDICES)

        # 1. Check if slot range is valid (within day, no breaks)
        if not is_slot_range_valid(start_slot_index, duration): continue

        # 2. Check ALL participating Department/Semester groups are free
        all_groups_free = True
        for group_key in participating_dept_sem_groups:
             if not check_availability(group_schedule, group_key, day_index, start_slot_index, duration):
                 all_groups_free = False
                 break
        if not all_groups_free: continue

        # 3. Check availability of ALL professors and find DISTINCT available rooms for each offering
        temp_prof_bookings = defaultdict(set) # Bookings for this attempt only
        assignments = {} # offering_key -> {'prof': prof, 'room': room_id}
        possible = True
        assigned_rooms_in_slot = set() # Keep track of rooms used by other offerings in this basket attempt

        for offering in offerings:
            offering_key = None
            prof = None
            enrollment = 0
            course_name_for_log = "Unknown Offering"
            offering_data = None # Store the original course/connector data for details

            if isinstance(offering, str): # It's a connector group key
                if offering not in connector_groups:
                    logging.warning(f"Basket {basket_key} references missing connector group {offering}. Skipping offering.")
                    possible = False; break # Skip this attempt if data inconsistent
                connector_info = connector_groups[offering]
                offering_key = offering
                prof = connector_info['professor']
                enrollment = connector_info['total_enrollment'] # Use max enrollment
                course_name_for_log = f"Connector {offering} ({connector_info['courses'][0]['Course Name']})"
                offering_data = connector_info # Store connector info

            elif isinstance(offering, dict): # Standalone course dict
                offering_key = offering['unique_id']
                prof = offering['Faculty']
                enrollment = offering['Enrolled_students']
                course_name_for_log = f"{offering['Course Name']} ({offering['Course Code']})"
                offering_data = offering # Store course info
            else: # Should not happen
                logging.error(f"Invalid offering type in basket {basket_key}: {offering}")
                possible = False; break

            # a. Check Professor availability (considering temp bookings for this slot)
            prof_free = True
            for i in range(duration):
                slot_idx = start_slot_index + i
                # Check real schedule AND temp bookings for this attempt
                if slot_idx in professor_schedule.get(prof, {}).get(day_index, set()) or \
                   slot_idx in temp_prof_bookings.get(prof, set()):
                    prof_free = False; break
            if not prof_free:
                logging.debug(f"Basket {basket_key} failed: Professor {prof} conflict for offering {course_name_for_log}")
                possible = False; break # Prof conflict within this basket attempt

            # b. Find a suitable, available, *distinct* room (assume LECTURE_ROOM for basket block)
            # Pass assigned_rooms_in_slot to exclude rooms already assigned *in this basket attempt*
            found_room = find_available_rooms(classrooms, 'LECTURE_ROOM', enrollment, day_index, start_slot_index, duration, classroom_schedule, num_rooms=1, excluded_rooms=assigned_rooms_in_slot)

            if found_room:
                room_id = found_room[0] # find_available_rooms returns a list, get the single room
                assignments[offering_key] = {'prof': prof, 'room': room_id}
                assigned_rooms_in_slot.add(room_id) # Add this room to the excluded list for the *next* offering in this basket
                # Temporarily mark prof as busy for subsequent checks *within this attempt*
                for i in range(duration):
                    slot_idx = start_slot_index + i
                    temp_prof_bookings[prof].add(slot_idx)
            else:
                logging.debug(f"Basket {basket_key} failed: No distinct room for {course_name_for_log} at {DAYS[day_index]} slot {start_slot_index} (Attempt {attempts})")
                possible = False; break # Cannot find room for this offering

        # 4. If all checks passed, book everything permanently
        if possible:
            logging.info(f"Successfully scheduled Basket {basket_key} at {DAYS[day_index]} Slot {start_slot_index}")

            # Book Dept/Sem Groups
            for group_key in participating_dept_sem_groups:
                 book_slot(group_schedule, group_key, day_index, start_slot_index, duration)

            # Book Professors, Rooms, and Update Timetables for each assignment
            for offering_key, assignment_info in assignments.items():
                prof = assignment_info['prof']
                room_id = assignment_info['room'] # This is a single room for the basket block

                  # Book prof and room permanently
                  book_slot(professor_schedule, prof, day_index, start_slot_index, duration)
                  book_slot(classroom_schedule, room_id, day_index, start_slot_index, duration)

                # Determine which courses/groups this assignment applies to and update timetable/status
                courses_to_update = []
                session_type = 'L' # Assume basket block is for Lecture
                course_name_tt = "Elective Offering" # Name for timetable
                course_code_tt = basket_key # Use basket key for the code in the timetable
                connector_id_for_log = None
                constituent_course_ids = [] # Unique IDs of courses affected by this scheduling

                if offering_key in connector_groups: # Connector group
                     connector_info = connector_groups[offering_key]
                     courses_to_update.extend(connector_info['courses']) # Update all courses in the connector
                    course_name_tt = f"BASKET: {connector_info['courses'][0]['Course Name']}" # Use first course name
                     course_code_tt = connector_info['courses'][0]['Course Code'] # Use first course code
                    connector_id_for_log = offering_key
                    constituent_course_ids = [c['unique_id'] for c in connector_info['courses']] # All courses in connector
                    # Mark connector group's L session as scheduled
                    connector_info['L_sessions_scheduled'] += 1

                else: # Standalone course (find the dict in the original offerings list)
                     found_course_dict = next((item for item in offerings if isinstance(item, dict) and item['unique_id'] == offering_key), None)
                    if found_course_dict:
                        courses_to_update.append(found_course_dict)
                        course_name_tt = f"BASKET: {found_course_dict['Course Name']}"
                        course_code_tt = found_course_dict['Course Code']
                        constituent_course_ids = [offering_key] # Just the standalone course
                        # Update status in the DataFrame for standalone elective L
                        course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == offering_key].tolist()
                        if course_idx:
                            course_schedule_status_df.loc[course_idx[0], f'{session_type}_scheduled'] += 1
                    else:
                        logging.error(f"Could not find offering data for key: {offering_key} in basket {basket_key}. Skipping timetable update for this offering.")
                        continue # Skip updating timetables for this offering if data not found


                # Update timetables for all affected department/semester groups
                for course_dict in courses_to_update:
                    dept_sem_key = f"{course_dict['Department']}_{course_dict['Semester']}"
                    if dept_sem_key in timetables:
                        timetable_entry_day = timetables[dept_sem_key][day_index]
                        for i in range(duration):
                            slot_idx = start_slot_index + i
                            is_first = (i == 0)
                            entry_type = session_type if is_first else 'CONT'
                            entry_span = duration if is_first else 0

                            timetable_entry_day[slot_idx] = {
                                'type': entry_type,
                                'course': course_name_tt,
                                'code': course_code_tt,
                                'rooms': [room_id], # Store room(s) as a list (single room for basket block)
                                'prof': prof,
                                'span': entry_span,
                                'connector': connector_id_for_log,
                                'basket': basket_key # Add basket info
                            }


            # Add to scheduled sessions list for the summary sheet
            # Avoid duplicating entries if a connector group represents multiple courses
            # Use a unique identifier for the scheduled session entry itself, e.g., basket_key + offering_key + day + start_slot
            session_list_key = f"{basket_key}_{offering_key}_{day_index}_{start_slot_index}"
            # Check if a session with this key is already added (could happen if a course belongs to multiple groups taking the same basket)
            if not any(item.get('SessionListKey') == session_list_key for item in scheduled_sessions_list):
                 scheduled_sessions_list.append({
                     'Course Code': course_code_tt,
                     'Course Name': course_name_tt,
                     'Type': session_type, # The type of the basket block (typically 'L')
                     'Day': DAYS[day_index],
                     'Start Time': ALL_SLOTS[start_slot_index]['start_time'].strftime('%H:%M'),
                     'End Time': ALL_SLOTS[start_slot_index + duration - 1]['end_time'].strftime('%H:%M'),
                     'Rooms': ", ".join([room_id]), # Single room for the basket block, joined for display
                     'Professor': prof,
                     # List all participating groups affected by this basket choice for this offering
                     'Groups': ", ".join(sorted(list(participating_dept_sem_groups))),
                     'Enrollment': enrollment, # Enrollment for this specific offering within the basket
                     'Connector': connector_id_for_log,
                     'Basket': basket_key,
                     'UniqueIDs': ", ".join(constituent_course_ids), # Store all unique IDs involved
                     'SessionListKey': session_list_key # Unique key for this list entry
                 })


            basket_info['scheduled'] = True
            basket_info['scheduled_day'] = day_index
            basket_info['scheduled_start_slot'] = start_slot_index
            # Store assignments in the basket info for easier retrieval if needed
            basket_info['scheduled_assignments'] = assignments

            scheduled = True
            return True # Successfully scheduled basket

    # --- Failed to schedule ---
    if not scheduled:
         logging.warning(f"Could not schedule Basket Group {basket_key} after {attempts} attempts.")
        # Mark constituent courses/connectors as having failed this session type (Lecture for basket)
        for offering in offerings:
            offering_key = offering if isinstance(offering, str) else offering['unique_id']
            if offering_key in connector_groups:
                 connector_info = connector_groups[offering_key]
                 connector_info['failed_sessions'].append({'type': 'L', 'count': 1}) # Assume failure for 1 session
            elif isinstance(offering, dict):
                 course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == offering_key].tolist()
                 if course_idx:
                    course_idx = course_idx[0]
                    failed_list = course_schedule_status_df.loc[course_idx, 'failed']
                    failed_list.append('L - Basket Failed') # Indicate failure due to basket
                    course_schedule_status_df.loc[course_idx, 'failed'] = failed_list # Ensure update is registered


    return False


def schedule_course_component(course_or_connector_info, session_type, duration, room_type_base,
                               classrooms, professor_schedule, classroom_schedule, group_schedule, timetables,
                               course_schedule_status_df, connector_groups, scheduled_sessions_list):
    """Attempt to schedule one session (L, T, or P) for a standalone course or a non-elective connector group."""

    # Determine if it's a connector group or a standalone course
    is_connector = isinstance(course_or_connector_info, dict) and 'connector_id' in course_or_connector_info

    # Get information needed for scheduling and logging
    if is_connector:
        group_info = course_or_connector_info
        connector_id = group_info['connector_id']

        # Skip if this connector group was already handled as part of a basket for this session type (L)
        # For L sessions, the basket scheduling takes precedence.
        if group_info.get('is_elective_connector', False) and session_type == 'L':
             # If the basket L session for this group was scheduled, we are done for L
             if group_info.get('L_sessions_scheduled', 0) > 0:
                  return True # Basket handled the L session

        prof = group_info['professor']
        enrollment = group_info['total_enrollment'] # Use max enrollment for room capacity
        participating_groups = group_info['participating_groups']
        log_name = f"Connector {connector_id} ({session_type})"
        unique_id_for_status = connector_id # Use connector_id for tracking status within connector_groups dict
        needed_key = f'{session_type}_sessions_needed'
        scheduled_key = f'{session_type}_sessions_scheduled'

        # Check if this session type is already fully scheduled for the group
        if group_info.get(scheduled_key, 0) >= group_info.get(needed_key, 0):
            return True # Already done for this type


    else: # Standalone course (provided as a dictionary)
        course = course_or_connector_info
        unique_id = course['unique_id']

        # Skip if handled by basket (assuming basket covers the L session for standalone electives)
        if course.get('is_elective', False) and session_type == 'L':
             # Check if the L session for this unique course is scheduled (by basket)
             course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == unique_id].tolist()
             if course_idx and course_schedule_status_df.loc[course_idx[0], 'L_scheduled'] > 0:
                  # Check if already fully scheduled
                  if course_schedule_status_df.loc[course_idx[0], f'{session_type}_scheduled'] >= course_schedule_status_df.loc[course_idx[0], f'{session_type}_sessions_needed']:
                       return True
             # If it's an elective L and L_scheduled is still 0, it means the basket scheduling failed for the parent basket.
             # We won't attempt to schedule it here as a standalone L. It's marked failed by basket scheduling.
             return False # Let basket scheduling handle elective L.


        prof = course['Faculty']
        enrollment = course['Enrolled_students']
        participating_groups = {f"{course['Department']}_{course['Semester']}"}
        log_name = f"{course['Course Name']} ({course['Course Code']}) {session_type}"
        unique_id_for_status = unique_id # Use unique_id for tracking status in the DataFrame
        needed_key = f'{session_type}_sessions_needed'
        scheduled_key = f'{session_type}_scheduled'

        # Check if already fully scheduled using the DataFrame status
        course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == unique_id_for_status].tolist()
        if course_idx and course_schedule_status_df.loc[course_idx[0], scheduled_key] >= course_schedule_status_df.loc[course_idx[0], needed_key]:
             return True


    # Determine required room type and number of rooms
    required_room_type = get_lab_type(course_or_connector_info['Department']) if room_type_base == 'LAB' else room_type_base
    num_rooms_needed = 2 if session_type == 'P' else 1 # P sessions need 2 rooms

    logging.debug(f"Attempting to schedule {session_type} ({num_rooms_needed} {'room' if num_rooms_needed == 1 else 'rooms'}) for {log_name} (Need {duration} slots)")

    attempts = 0
    scheduled = False
    while not scheduled and attempts < MAX_SCHEDULING_ATTEMPTS_SESSION:
        attempts += 1
        day_index = random.randrange(len(DAYS))
        # Pick a random valid starting slot index from schedulable slots
        if not SCHEDULABLE_SLOT_INDICES:
            logging.error("No schedulable slots available!")
            break
        start_slot_index = random.choice(SCHEDULABLE_SLOT_INDICES)

        # 1. Check slot validity (within day, no breaks)
        if not is_slot_range_valid(start_slot_index, duration): continue

        # 2. Check Professor Availability
        if not check_availability(professor_schedule, prof, day_index, start_slot_index, duration): continue

        # 3. Check Group(s) Availability
        all_groups_free = all(check_availability(group_schedule, g_key, day_index, start_slot_index, duration) for g_key in participating_groups)
        if not all_groups_free: continue

        # 4. Find Available Classroom(s)
        found_rooms = find_available_rooms(classrooms, required_room_type, enrollment, day_index, start_slot_index, duration, classroom_schedule, num_rooms=num_rooms_needed)

        if found_rooms and len(found_rooms) == num_rooms_needed:
             # --- Success! Book the slot ---
             logging.info(f"Scheduled: {session_type} ({','.join(found_rooms)}) for {log_name} at {DAYS[day_index]} {ALL_SLOTS[start_slot_index]['label']} by {prof}")

             # Book resources
             book_slot(professor_schedule, prof, day_index, start_slot_index, duration)
            # Book each found room
            for room_id in found_rooms:
                 book_slot(classroom_schedule, room_id, day_index, start_slot_index, duration)
            # Book all participating groups
             for g_key in participating_groups:
                 book_slot(group_schedule, g_key, day_index, start_slot_index, duration)


             # Update timetable(s) and status
            connector_id_for_log = None
            course_name_tt = None # Course name for timetable
            course_code_tt = None # Course code for timetable
            constituent_course_ids = [] # Unique IDs of courses affected

            if is_connector:
                group_info[scheduled_key] += 1 # Increment connector scheduled count
                connector_id_for_log = connector_id
                # Use info from the first course in the connector group for timetable entry
                first_course_in_group = group_info['courses'][0]
                course_name_tt = f"CONN: {first_course_in_group['Course Name']}"
                course_code_tt = first_course_in_group['Course Code']
                constituent_course_ids = [c['unique_id'] for c in group_info['courses']] # All courses in connector

                # Update status for all constituent courses in the DataFrame
                for c_dict in group_info['courses']:
                    if c_dict['unique_id'] in course_schedule_status_df['unique_id'].values:
                        course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == c_dict['unique_id']].tolist()
                        if course_idx:
                             # Ensure we don't double count if a session is only needed once for the group
                             # Check if the count for this specific course is less than needed (should be 0 or 1 increment per scheduled session)
                             if course_schedule_status_df.loc[course_idx[0], scheduled_key] < c_dict[f'{session_type}_sessions_needed']:
                                 course_schedule_status_df.loc[course_idx[0], scheduled_key] += 1

            else: # Standalone course
                course = course_or_connector_info
                course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == unique_id_for_status].tolist()
                if course_idx:
                    course_schedule_status_df.loc[course_idx[0], scheduled_key] += 1
                course_name_tt = course['Course Name']
                course_code_tt = course['Course Code']
                constituent_course_ids = [unique_id_for_status] # Just the standalone course


             # Update timetables for all relevant groups
             for g_key in participating_groups:
                 if g_key in timetables:
                     timetable_entry_day = timetables[g_key][day_index]
                     for i in range(duration):
                         slot_idx = start_slot_index + i
                         is_first = (i == 0)
                         entry_type = session_type if is_first else 'CONT'
                         entry_span = duration if is_first else 0
                         timetable_entry_day[slot_idx] = {
                           'type': entry_type,
                           'course': course_name_tt,
                           'code': course_code_tt,
                           'rooms': found_rooms, # Store list of rooms
                           'prof': prof,
                           'span': entry_span,
                           'connector': connector_id_for_log,
                           'basket': None # Not scheduled as part of a basket directly
                       }

            # Add to scheduled sessions list for the summary sheet
            # Create a unique key for this session entry to avoid duplicates in the list
            session_list_key = f"{','.join(sorted(constituent_course_ids))}_{session_type}_{day_index}_{start_slot_index}_{','.join(sorted(found_rooms))}"
            if not any(item.get('SessionListKey') == session_list_key for item in scheduled_sessions_list):
                 scheduled_sessions_list.append({
                     'Course Code': course_code_tt, # Use code from timetable entry
                     'Course Name': course_name_tt, # Use name from timetable entry
                     'Type': session_type,
                     'Day': DAYS[day_index],
                     'Start Time': ALL_SLOTS[start_slot_index]['start_time'].strftime('%H:%M'),
                     'End Time': ALL_SLOTS[start_slot_index + duration - 1]['end_time'].strftime('%H:%M'),
                     'Rooms': ", ".join(found_rooms), # Join rooms for the summary sheet
                     'Professor': prof,
                     'Groups': ", ".join(sorted(list(participating_groups))),
                     'Enrollment': enrollment, # Enrollment used for room capacity
                     'Connector': connector_id_for_log,
                     'Basket': None, # Not scheduled as part of a basket block
                     'UniqueIDs': ", ".join(constituent_course_ids), # Store all unique IDs involved
                     'SessionListKey': session_list_key # Unique key for this list entry
                 })


            scheduled = True
            return True

    # --- Failed to schedule after attempts ---
    if not scheduled:
        logging.warning(f"Could not schedule {session_type} ({num_rooms_needed} rooms) for {log_name} after {attempts} attempts.")
        if is_connector:
            # Find the count already failed for this type to avoid duplicates if retry logic were implemented
            existing_failed_count = sum(item['count'] for item in group_info['failed_sessions'] if item['type'] == session_type)
            if existing_failed_count < group_info.get(needed_key, 0) - group_info.get(scheduled_key, 0):
                 group_info['failed_sessions'].append({'type': session_type, 'count': 1})
                 logging.warning(f"Marked one {session_type} session as failed for Connector {connector_id}.")
            else:
                 logging.debug(f"Connector {connector_id} {session_type} already marked as fully failed or scheduled.")

        else: # Standalone course
            course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == unique_id_for_status].tolist()
            if course_idx:
                course_idx = course_idx[0]
                failed_list = course_schedule_status_df.loc[course_idx, 'failed']
                # Only add if not already marked failed for this type
                if session_type not in failed_list:
                    failed_list.append(session_type)
                    course_schedule_status_df.loc[course_idx, 'failed'] = failed_list # Ensure update is registered
                    logging.warning(f"Marked {session_type} as failed for {unique_id_for_status}.")
                else:
                    logging.debug(f"{session_type} for {unique_id_for_status} already marked as failed.")


    return False

# --- Output Generation ---
def export_to_excel(timetables, course_schedule_status_df, connector_groups, basket_groups, scheduled_sessions_list, filename="timetable_output.xlsx"):
    """Export the generated timetables, summary, and status to an Excel file."""
    wb = Workbook()
    wb.remove(wb.active) # Remove default sheet

    header_font = Font(bold=True)
    center_alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    vertical_alignment = Alignment(vertical='center', wrap_text=True)

    fill_lecture = PatternFill(start_color="ADD8E6", end_color="ADD8E6", fill_type="solid") # Light Blue
    fill_tutorial = PatternFill(start_color="90EE90", end_color="90EE90", fill_type="solid") # Light Green
    fill_lab = PatternFill(start_color="FFB6C1", end_color="FFB6C1", fill_type="solid") # Light Pink
    fill_basket = PatternFill(start_color="FFD700", end_color="FFD700", fill_type="solid") # Gold
    fill_connector = PatternFill(start_color="D3D3D3", end_color="D3D3D3", fill_type="solid") # Light Grey
    fill_cont = PatternFill(start_color="F5F5DC", end_color="F5F5DC", fill_type="solid") # Beige
    fill_break = PatternFill(start_color="CCCCCC", end_color="CCCCCC", fill_type="solid") # Grey for breaks


    # --- Create Timetable Sheets ---
    for group_key, timetable in timetables.items():
        ws = wb.create_sheet(title=group_key)

        # Write headers
        ws.cell(row=1, column=1, value="Time Slot").font = header_font
        ws.column_dimensions['A'].width = 15 # Adjust width for time slots
        for i, day in enumerate(DAYS):
            col_letter = get_column_letter(i + 2) # Start from column B
            ws.cell(row=1, column=i + 2, value=day).font = header_font
            ws.column_dimensions[col_letter].width = 25 # Adjust width for course info

        # Write time slots and schedule entries
        for slot_index, slot_info in enumerate(ALL_SLOTS):
            ws.cell(row=slot_index + 2, column=1, value=slot_info['label']).alignment = vertical_alignment

            # Apply break color to the time slot column cell as well
            if slot_info['is_break']:
                 ws.cell(row=slot_index + 2, column=1).fill = fill_break

            for day_index, day_schedule in enumerate(timetable):
                entry = day_schedule[slot_index]
                cell_value = ""
                cell_fill = None

                # Apply break color to all cells in break rows
                if slot_info['is_break']:
                    cell = ws.cell(row=slot_index + 2, column=day_index + 2, value="BREAK")
                    cell.alignment = center_alignment
                    cell.font = Font(italic=True)
                    cell.fill = fill_break
                    continue


                if entry['type']:
                    # Determine fill color based on type/properties
                    if entry['basket']: # Prioritize basket color if part of a basket block
                        cell_fill = fill_basket
                    elif entry['connector']: # Prioritize connector color if part of a connector group
                        cell_fill = fill_connector
                    elif entry['type'] == 'L':
                        cell_fill = fill_lecture
                    elif entry['type'] == 'T':
                        cell_fill = fill_tutorial
                    elif entry['type'] == 'P':
                        cell_fill = fill_lab
                    elif entry['type'] == 'CONT':
                        cell_fill = fill_cont


                    if entry['span'] > 0: # This is the start of a session
                        rooms_str = ", ".join(entry.get('rooms', [])) if entry.get('rooms') else "N/A"
                        cell_value = f"{entry['course']} ({entry['code']})\nProf: {entry['prof']}\nRoom: {rooms_str}"
                        # Add basket/connector info to the cell text if applicable
                        if entry['basket']:
                            cell_value += f"\nBasket: {entry['basket']}"
                        elif entry['connector']:
                             cell_value += f"\nConnector: {entry['connector']}"


                        cell = ws.cell(row=slot_index + 2, column=day_index + 2, value=cell_value)
                        cell.alignment = center_alignment
                        if cell_fill: cell.fill = cell_fill

                        # Merge cells for the duration of the session
                        if entry['span'] > 1:
                            ws.merge_cells(start_row=slot_index + 2, start_column=day_index + 2,
                                          end_row=slot_index + 2 + entry['span'] - 1, end_column=day_index + 2)

                    else: # Continuation of a session ('CONT')
                        # Cell value is empty, it's covered by the merged cell above.
                        # We still set the fill color for consistency, although merge might override.
                        cell = ws.cell(row=slot_index + 2, column=day_index + 2)
                        if cell_fill: cell.fill = cell_fill
                        pass # No value needed for continuation cells


    # --- Create Scheduled Classes List Sheet ---
    ws_list = wb.create_sheet(title="Scheduled Class List")
    # Define headers for the list sheet
    list_headers = ["Course Code", "Course Name", "Type", "Day", "Start Time", "End Time", "Rooms", "Professor", "Groups", "Enrollment", "Connector", "Basket"]
    for col_num, header in enumerate(list_headers, 1):
        ws_list.cell(row=1, column=col_num, value=header).font = header_font
        ws_list.column_dimensions[get_column_letter(col_num)].width = 20 # Adjust width

    row_num = 2
    # Sort scheduled_sessions_list for better readability (e.g., by Day, Start Time, Group)
    scheduled_sessions_list.sort(key=lambda x: (DAYS.index(x['Day']), x['Start Time'], x['Groups']))

    for session in scheduled_sessions_list:
        col_num = 1
        for header in list_headers:
            # Use .get() with a default value to handle potential missing keys gracefully
            cell_value = session.get(header, "")
            ws_list.cell(row=row_num, column=col_num, value=cell_value).alignment = vertical_alignment
            col_num += 1
        row_num += 1


    # --- Create Scheduling Status Sheet (for individual courses) ---
    ws_status = wb.create_sheet(title="Course Scheduling Status")
    status_headers = ["Department", "Semester", "Course Code", "Course Name", "Faculty",
                      "L Needed", "L Scheduled", "T Needed", "T Scheduled", "P Needed", "P Scheduled", "Failed Sessions"]
    for col_num, header in enumerate(status_headers, 1):
        ws_status.cell(row=1, column=col_num, value=header).font = header_font
        ws_status.column_dimensions[get_column_letter(col_num)].width = 15

    row_num = 2
    # Sort status by Department and Semester
    course_schedule_status_df_sorted = course_schedule_status_df.sort_values(by=['Department', 'Semester', 'Course Code'])

    for index, row in course_schedule_status_df_sorted.iterrows():
        ws_status.cell(row=row_num, column=1, value=row['Department']).alignment = vertical_alignment
        ws_status.cell(row=row_num, column=2, value=row['Semester']).alignment = vertical_alignment
        ws_status.cell(row=row_num, column=3, value=row['Course Code']).alignment = vertical_alignment
        ws_status.cell(row=row_num, column=4, value=row['Course Name']).alignment = vertical_alignment
        ws_status.cell(row=row_num, column=5, value=row['Faculty']).alignment = vertical_alignment
        ws_status.cell(row=row_num, column=6, value=row['L_sessions_needed']).alignment = center_alignment
        ws_status.cell(row=row_num, column=7, value=row['L_scheduled']).alignment = center_alignment
        ws_status.cell(row=row_num, column=8, value=row['T_sessions_needed']).alignment = center_alignment
        ws_status.cell(row=row_num, column=9, value=row['T_scheduled']).alignment = center_alignment
        ws_status.cell(row=row_num, column=10, value=row['P_sessions_needed']).alignment = center_alignment
        ws_status.cell(row=row_num, column=11, value=row['P_scheduled']).alignment = center_alignment
        # Convert list of failed sessions to a string
        ws_status.cell(row=row_num, column=12, value=", ".join(row['failed'])).alignment = vertical_alignment
        row_num += 1

    # Add Connector Status Sheet
    ws_conn_status = wb.create_sheet(title="Connector Status")
    conn_status_headers = ["Connector ID", "Professor", "Participating Groups", "L Needed", "L Scheduled", "T Needed", "T Scheduled", "P Needed", "P Scheduled", "Failed Sessions", "Is Elective", "Basket Code"]
    for col_num, header in enumerate(conn_status_headers, 1):
        ws_conn_status.cell(row=1, column=col_num, value=header).font = header_font
        ws_conn_status.column_dimensions[get_column_letter(col_num)].width = 15

    row_num = 2
    # Sort connectors by ID
    sorted_connector_groups = sorted(connector_groups.values(), key=lambda x: x.get('connector_id', ''))

    for conn_info in sorted_connector_groups:
        ws_conn_status.cell(row=row_num, column=1, value=conn_info.get('connector_id', '')).alignment = vertical_alignment
        ws_conn_status.cell(row=row_num, column=2, value=conn_info.get('professor', '')).alignment = vertical_alignment
        ws_conn_status.cell(row=row_num, column=3, value=", ".join(sorted(list(conn_info.get('participating_groups', set()))))).alignment = vertical_alignment
        ws_conn_status.cell(row=row_num, column=4, value=conn_info.get('L_sessions_needed', 0)).alignment = center_alignment
        ws_conn_status.cell(row=row_num, column=5, value=conn_info.get('L_sessions_scheduled', 0)).alignment = center_alignment
        ws_conn_status.cell(row=row_num, column=6, value=conn_info.get('T_sessions_needed', 0)).alignment = center_alignment
        ws_conn_status.cell(row=row_num, column=7, value=conn_info.get('T_sessions_scheduled', 0)).alignment = center_alignment
        ws_conn_status.cell(row=row_num, column=8, value=conn_info.get('P_sessions_needed', 0)).alignment = center_alignment
        ws_conn_status.cell(row=row_num, column=9, value=conn_info.get('P_sessions_scheduled', 0)).alignment = center_alignment
        # Convert failed sessions list to a string
        failed_list_str = ", ".join([f"{f.get('type', 'Unknown')}({f.get('count', 1)})" for f in conn_info.get('failed_sessions', [])])
        ws_conn_status.cell(row=row_num, column=10, value=failed_list_str).alignment = vertical_alignment
        ws_conn_status.cell(row=row_num, column=11, value=conn_info.get('is_elective_connector', False)).alignment = center_alignment
        ws_conn_status.cell(row=row_num, column=12, value=conn_info.get('basket_code', '') if conn_info.get('basket_code') else "").alignment = center_alignment

        row_num += 1

    # Add Basket Status Sheet
    ws_basket_status = wb.create_sheet(title="Basket Status")
    basket_status_headers = ["Basket Key", "Basket Code", "Semester Group", "Participating Groups", "Number of Offerings", "Scheduled", "Scheduled Day", "Scheduled Start Slot"]
    for col_num, header in enumerate(basket_status_headers, 1):
        ws_basket_status.cell(row=1, column=col_num, value=header).font = header_font
        ws_basket_status.column_dimensions[get_column_letter(col_num)].width = 20

    row_num = 2
    # Sort baskets by key
    sorted_basket_groups = sorted(basket_groups.values(), key=lambda x: x.get('basket_code', ''))

    for basket_info in sorted_basket_groups:
        ws_basket_status.cell(row=row_num, column=1, value=f"{basket_info.get('semester_logic_group', '')}_{basket_info.get('basket_code', '')}").alignment = vertical_alignment
        ws_basket_status.cell(row=row_num, column=2, value=basket_info.get('basket_code', '')).alignment = vertical_alignment
        ws_basket_status.cell(row=row_num, column=3, value=basket_info.get('semester_logic_group', '')).alignment = vertical_alignment
        ws_basket_status.cell(row=row_num, column=4, value=", ".join(sorted(list(basket_info.get('participating_groups', set()))))).alignment = vertical_alignment
        ws_basket_status.cell(row=row_num, column=5, value=len(basket_info.get('offerings', []))).alignment = center_alignment
        ws_basket_status.cell(row=row_num, column=6, value=basket_info.get('scheduled', False)).alignment = center_alignment
        ws_basket_status.cell(row=row_num, column=7, value=DAYS[basket_info['scheduled_day']] if basket_info.get('scheduled_day') is not None else "N/A").alignment = center_alignment
        ws_basket_status.cell(row=row_num, column=8, value=ALL_SLOTS[basket_info['scheduled_start_slot']]['label'] if basket_info.get('scheduled_start_slot') is not None else "N/A").alignment = center_alignment

        row_num += 1


    # Save the workbook
    try:
        wb.save(filename)
        logging.info(f"Timetable successfully exported to {filename}")
    except Exception as e:
        logging.error(f"Error saving Excel file: {e}")


# --- Main Scheduling Logic ---
def generate_timetable():
    logging.info("Starting timetable generation...")

    # Create dummy CSV files if they don't exist for testing
    if not os.path.exists('Classrooms.csv'):
        logging.warning("Classrooms.csv not found. Creating a dummy file.")
        dummy_classrooms = pd.DataFrame({
            'Classroom': [f'CR{i}' for i in range(1, 11)] + [f'LAB{i}' for i in range(1, 7)],
            'Capacity': [120, 80, 60, 60, 50, 50, 40, 40, 30, 30, 50, 50, 45, 45, 40, 40],
            'Type': ['LECTURE_ROOM'] * 8 + ['TUTORIAL_ROOM'] * 2 + ['COMPUTER_LAB'] * 3 + ['HARDWARE_LAB'] * 3
        })
        dummy_classrooms.to_csv('Classrooms.csv', index=False)

    if not os.path.exists('Semester_courses.csv'):
        logging.warning("Semester_courses.csv not found. Creating a dummy file.")
        dummy_courses = pd.DataFrame({
            'Department': ['CSE', 'CSE', 'CSE', 'CSE', 'CSE', 'ECE', 'ECE', 'DSAI', 'DSAI', 'CSE', 'ECE', 'DSAI'],
            'Semester': ['6', '6', '6', '6', '6', '6', '6', '6', '6', '4', '4', '4'],
            'Course Code': ['CS601', 'CS602', 'CS603', 'CS604', 'B1(CSB101)', 'EC601', 'EC602', 'DS601', 'B1(DSB102)', 'CS401', 'EC401', 'DS401'],
            'Course Name': ['Advanced Algo', 'ML', 'Network Security', 'Cloud Computing', 'Basket 1 Elective A', 'Digital Comm', 'VLSI Design', 'Data Mining', 'Basket 1 Elective B', 'OS', 'Signals', 'AI Basics'],
            'Faculty': ['Dr. Smith', 'Dr. Jones', 'Dr. Lee', 'Dr. Gupta', 'Dr. Jones', 'Dr. Khan', 'Dr. Reddy', 'Dr. Sharma', 'Dr. Sharma', 'Dr. Smith', 'Dr. Khan', 'Dr. Sharma'],
            'Enrolled_students': [80, 90, 65, 70, 50, 75, 60, 55, 40, 70, 60, 50],
            'L': [3, 3, 2, 2, 3, 3, 2, 3, 3, 3, 3, 3],
            'T': [1, 1, 1, 0, 0, 1, 0, 1, 0, 1, 1, 1],
            'P': [0, 0, 1, 1, 0, 0, 1, 0, 0, 1, 1, 1],
            'class_connector': ['', '', '', '', 'B1_Group', '', '', '', 'B1_Group', '', '', '']
        })
        dummy_courses.to_csv('Semester_courses.csv', index=False)


    # Load data
    logging.info("Loading classroom data...")
    classrooms = load_classroom_data('Classrooms.csv')
    logging.info("Loading course data...")
    courses_df = load_course_data('Semester_courses.csv')

    # Identify unique departments and semesters
    departments = courses_df['Department'].unique().tolist()
    semesters_by_dept = courses_df.groupby('Department')['Semester'].unique().to_dict()

    # Initialize schedules and timetable structure
    professor_schedule = {} # {prof: {day_index: set(busy_slot_indices)}}
    classroom_schedule = {} # {room_id: {day_index: set(busy_slot_indices)}}
    group_schedule = {} # {dept_sem_key: {day_index: set(busy_slot_indices)}}
    timetables = initialize_timetable(departments, semesters_by_dept)

    # Use the DataFrame itself to track scheduling status for individual courses
    course_schedule_status_df = courses_df.copy()

    # Identify and preprocess connector and basket groups
    logging.info("Identifying connector groups...")
    connector_groups = build_connector_groups(courses_df)
    logging.info(f"Found {len(connector_groups)} connector groups.")

    logging.info("Identifying basket groups...")
    basket_groups = group_courses_by_basket(courses_df, connector_groups)
    logging.info(f"Found {len(basket_groups)} basket groups.")


    # Keep track of scheduled sessions for the summary list sheet
    scheduled_sessions_list = []

    # --- Scheduling Strategy ---
    # 1. Schedule Basket Groups first (common elective block - typically Lecture)
    logging.info("\n--- Scheduling Basket Groups ---")
    for basket_key, basket_info in basket_groups.items():
         schedule_basket_group(basket_key, basket_info, classrooms, connector_groups,
                               professor_schedule, classroom_schedule, group_schedule, timetables, course_schedule_status_df, scheduled_sessions_list)


    # 2. Schedule regular sessions (L, P, T) for standalone courses and non-elective connectors
    # Prioritize Labs (P) first as they need specific rooms and multiple rooms
    # Then Lectures (L), then Tutorials (T)

    sessions_to_schedule = []

    # Add standalone courses (excluding elective L sessions, handled by baskets)
    # Iterate through the status DataFrame to get current needs
    for index, course_row in course_schedule_status_df.iterrows():
        course_dict = course_row.to_dict() # Convert row to dict for schedule function

        # Skip if this course is part of a connector group (will be handled via connector entry)
        if course_dict.get('class_connector'):
            continue

        # Schedule L (if not an elective, or if elective L failed basket scheduling - although current logic relies on basket L)
        if not course_dict.get('is_elective', False) and course_dict['L_sessions_needed'] > course_dict['L_scheduled']:
             sessions_to_schedule.append({'info': course_dict, 'type': 'L', 'needed': course_dict['L_sessions_needed'], 'scheduled': course_dict['L_scheduled'], 'duration': LECTURE_SLOTS, 'room_type': 'LECTURE_ROOM'})

        # Schedule T (for all courses, including electives)
        if course_dict['T_sessions_needed'] > course_dict['T_scheduled']:
            sessions_to_schedule.append({'info': course_dict, 'type': 'T', 'needed': course_dict['T_sessions_needed'], 'scheduled': course_dict['T_scheduled'], 'duration': TUTORIAL_SLOTS, 'room_type': 'LECTURE_ROOM'}) # Assume T in lecture rooms

        # Schedule P (for all courses, including electives)
        if course_dict['P_sessions_needed'] > course_dict['P_scheduled']:
            sessions_to_schedule.append({'info': course_dict, 'type': 'P', 'needed': course_dict['P_sessions_needed'], 'scheduled': course_dict['P_scheduled'], 'duration': LAB_SLOTS, 'room_type': 'LAB'})


    # Add non-elective connector groups (schedule L, T, P for the group as a whole)
    for connector_id, group_info in connector_groups.items():
        if not group_info.get('is_elective_connector', False): # Only non-elective connectors
            if group_info['L_sessions_needed'] > group_info['L_sessions_scheduled']:
                 sessions_to_schedule.append({'info': group_info, 'type': 'L', 'needed': group_info['L_sessions_needed'], 'scheduled': group_info['L_sessions_scheduled'], 'duration': LECTURE_SLOTS, 'room_type': 'LECTURE_ROOM'})
            if group_info['T_sessions_needed'] > group_info['T_sessions_scheduled']:
                 sessions_to_schedule.append({'info': group_info, 'type': 'T', 'needed': group_info['T_sessions_needed'], 'scheduled': group_info['T_sessions_scheduled'], 'duration': TUTORIAL_SLOTS, 'room_type': 'LECTURE_ROOM'}) # Assume T in lecture rooms
            if group_info['P_sessions_needed'] > group_info['P_sessions_scheduled']:
                 sessions_to_schedule.append({'info': group_info, 'type': 'P', 'needed': group_info['P_sessions_needed'], 'scheduled': group_info['P_sessions_scheduled'], 'duration': LAB_SLOTS, 'room_type': 'LAB'})


    # Prioritize scheduling: P > L > T
    sessions_to_schedule.sort(key=lambda x: (0 if x['type'] == 'P' else (1 if x['type'] == 'L' else 2)))


    logging.info("\n--- Scheduling Regular Sessions (P, L, T) ---")
    # Iterate and schedule sessions until all needed sessions are scheduled or attempts are exhausted
    # We loop through the types until no more sessions of that type can be scheduled
    # This is a simple approach; a more advanced scheduler might re-attempt failed sessions later
    something_scheduled_in_pass = True
    while something_scheduled_in_pass:
        something_scheduled_in_pass = False
        remaining_sessions_in_pass = []

        for session in sessions_to_schedule:
             info = session['info']
             session_type = session['type']
             needed_count = session['needed']
             scheduled_count = session['scheduled'] # Use the count from the session dict for loop control
             duration = session['duration']
             room_type_base = session['room_type']

             # Re-fetch current scheduled count from the source of truth (DataFrame or connector_groups)
             if isinstance(info, dict) and 'connector_id' in info:
                 current_scheduled_count = connector_groups[info['connector_id']].get(f'{session_type}_sessions_scheduled', 0)
             else:
                 course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == info['unique_id']].tolist()
                 if course_idx:
                      current_scheduled_count = course_schedule_status_df.loc[course_idx[0], f'{session_type}_scheduled']
                 else:
                      # Should not happen, but handle defensively
                      logging.error(f"Could not find status for {info.get('unique_id', info.get('connector_id', 'Unknown'))}. Skipping session.")
                      continue

             if current_scheduled_count < needed_count:
                 # Attempt to schedule one session
                 success = schedule_course_component(info, session_type, duration, room_type_base,
                                                     classrooms, professor_schedule, classroom_schedule, group_schedule, timetables,
                                                     course_schedule_status_df, connector_groups, scheduled_sessions_list)
                 if success:
                      something_scheduled_in_pass = True # Mark that progress was made
                      # Update the scheduled count in the session dict for the next iteration of this pass (or just rely on re-fetch)
                      # Re-fetching is more reliable:
                      if isinstance(info, dict) and 'connector_id' in info:
                            session['scheduled'] = connector_groups[info['connector_id']].get(f'{session_type}_sessions_scheduled', 0)
                      else:
                            course_idx = course_schedule_status_df.index[course_schedule_status_df['unique_id'] == info['unique_id']].tolist()
                            if course_idx:
                                 session['scheduled'] = course_schedule_status_df.loc[course_idx[0], f'{session_type}_scheduled']

                 # Add session back to list if more are needed
                 if session['scheduled'] < needed_count:
                     remaining_sessions_in_pass.append(session)
             # If already scheduled >= needed, this session is complete for this type, don't add back

        sessions_to_schedule = remaining_sessions_in_pass # Process remaining needed sessions in the next pass


    # --- Final Output ---
    logging.info("\nScheduling complete. Generating output.")
    export_to_excel(timetables, course_schedule_status_df, connector_groups, basket_groups, scheduled_sessions_list)
    logging.info("Timetable generation finished.")


# --- Entry Point ---
if __name__ == "__main__":
    generate_timetable()